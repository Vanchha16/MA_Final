"""Turn the requirement docs (Markdown) into Word files for submission.

pandoc (from pypandoc_binary) converts; python-docx then adds table borders,
header shading, the brand font and fits the diagrams on the page.
Usage: python make_docx.py <docs/requirements dir>
"""
import sys
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

BRAND = RGBColor(0x02, 0x84, 0xC7)
MAX_W, MAX_H = Inches(6.5), Inches(7.9)


def borders(table):
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "94A3B8")
        b.append(e)
    tbl_pr.append(b)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:fill"), fill)
    tc_pr.append(s)


def widen_id_column(table):
    """Tables size their columns automatically; keep IDs like "FR-U01" on one line."""
    first = [r.cells[0].text.strip() for r in table.rows[1:]]
    if not first or not all(len(t) <= 8 and "-" in t for t in first):
        return
    for row in table.rows:
        tc_pr = row.cells[0]._tc.get_or_add_tcPr()
        tc_pr.append(OxmlElement("w:noWrap"))


def keep_heading_with_next(doc):
    for p in doc.paragraphs:
        if p.style is not None and (p.style.name or "").lower().startswith("heading"):
            p.paragraph_format.keep_with_next = True


def polish(path):
    doc = Document(path)
    for style in doc.styles:
        if style.type == 1 and style.font is not None:  # paragraph styles
            style.font.name = "Roboto"
    sizes = {"title": 22, "heading 1": 16, "heading 2": 14, "heading 3": 12}
    for st in doc.styles:
        key = (st.name or "").lower()
        if key in sizes:
            st.font.color.rgb = BRAND
            st.font.size = Pt(sizes[key])
            st.font.bold = True
        elif key in ("normal", "body text", "first paragraph", "compact"):
            st.font.size = Pt(10.5)
    keep_heading_with_next(doc)
    for t in doc.tables:
        if all(not c.text.strip() for c in t.rows[0].cells):
            # key / value tables ("| | |" in Markdown): drop the empty header row
            t._tbl.remove(t.rows[0]._tr)
            borders(t)
            for row in t.rows:
                shade(row.cells[0], "E0F2FE")
                for p in row.cells[0].paragraphs:
                    for r in p.runs:
                        r.font.bold = True
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(9.5)
            continue
        widen_id_column(t)
        borders(t)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell in t.rows[0].cells:
            shade(cell, "E0F2FE")
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(9.5)
    for shape in doc.inline_shapes:
        w, h = shape.width, shape.height
        k = min(MAX_W / w, MAX_H / h, 1)
        shape.width, shape.height = Emu(int(w * k)), Emu(int(h * k))
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(0.9)
    doc.save(path)


if __name__ == "__main__":
  SRC = Path(sys.argv[1])
  for md in ("REQUIREMENTS.md", "USE_CASES.md"):
    out = SRC / md.replace(".md", ".docx")
    pypandoc.convert_file(
        str(SRC / md), "docx", outputfile=str(out),
        extra_args=[f"--resource-path={SRC}", "--from=gfm"],
    )
    polish(out)
    print("wrote", out)
